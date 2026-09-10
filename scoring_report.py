"""Generates a single Excel workbook with the final, all-rounds-combined
standings for every event -- meant for once a rodeo is over, to hand off
to a committee/association or archive alongside the exported XML files.

Uses openpyxl (already a dependency for reading Excel draw sheets on
import -- see importer.py) rather than adding a PDF-generation library,
so there's nothing new to install on any OS.

Ranking rules exactly match aggregate_leaderboard.xml (see
xml_export.export_aggregate_leaderboard): scores/times are summed across
every round in the event, more completed head always ranks above fewer
regardless of raw total, and only entries with status == 'scored' count
toward a rank at all. Anyone entered in the event but without a single
scored run (DNS, DQ, still pending, etc.) is listed separately per event
rather than silently dropped, since a committee handling payouts usually
still wants that on record.
"""
import datetime
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from database import get_conn, get_setting
import scoring
from xml_export import _team_names

TITLE_FONT = Font(name="Arial", bold=True, size=16)
SUBTITLE_FONT = Font(name="Arial", size=11, italic=True, color="666666")
SECTION_FONT = Font(name="Arial", bold=True, size=12)
HEADER_FONT = Font(name="Arial", bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5496")
BODY_FONT = Font(name="Arial", size=11)
CENTER = Alignment(horizontal="center")

# Excel sheet name limits: 31 chars max, and none of : \ / ? * [ ]
_SHEET_NAME_BAD_CHARS = re.compile(r'[:\\/?*\[\]]')


def _safe_sheet_title(name, used):
    """A sanitized, unique-within-this-workbook sheet title. Two events
    could plausibly share a name (or both sanitize down to the same
    thing), so this disambiguates with a "(2)", "(3)", ... suffix rather
    than letting openpyxl raise on a duplicate."""
    cleaned = _SHEET_NAME_BAD_CHARS.sub("", name).strip()[:31] or "Event"
    candidate = cleaned
    n = 2
    while candidate in used:
        suffix = f" ({n})"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate)
    return candidate


def _event_rows(conn, event):
    """(ranked scored rows, no-qualifying-score rows) for one event,
    aggregated across every round -- same query shape as
    export_aggregate_leaderboard, just not limited to the live event."""
    round_ids = [r["id"] for r in conn.execute(
        "SELECT id FROM rounds WHERE event_id = ?", (event["id"],)
    )]
    if not round_ids:
        return [], []

    placeholders = ",".join("?" * len(round_ids))
    scored = conn.execute(
        f"""
        SELECT c.id, c.name, c.hometown, c.sponsor,
               COALESCE(pc.name, c.partner) as partner,
               SUM(e.score_value) as total, COUNT(e.id) as head
        FROM entries e
        JOIN competitors c ON e.competitor_id = c.id
        LEFT JOIN competitors pc ON c.partner_id = pc.id
        WHERE e.round_id IN ({placeholders}) AND e.status = 'scored'
        GROUP BY c.id
        """,
        round_ids,
    ).fetchall()

    no_score = conn.execute(
        f"""
        SELECT c.name, COALESCE(pc.name, c.partner) as partner,
               GROUP_CONCAT(DISTINCT e.status) as statuses
        FROM entries e
        JOIN competitors c ON e.competitor_id = c.id
        LEFT JOIN competitors pc ON c.partner_id = pc.id
        WHERE e.round_id IN ({placeholders})
        GROUP BY c.id
        HAVING SUM(CASE WHEN e.status = 'scored' THEN 1 ELSE 0 END) = 0
        """,
        round_ids,
    ).fetchall()

    scored = [dict(r) for r in scored]
    no_score = [dict(r) for r in no_score]

    # Same reasoning as the XML exports: a partner_id is a property of
    # the person, not this event, so only show it when this event is
    # actually marked as a team event.
    if not event["is_team"]:
        for r in scored:
            r["partner"] = ""
        for r in no_score:
            r["partner"] = ""

    scoring_type = event["scoring_type"]
    if scoring.is_timed(scoring_type):
        scored.sort(key=lambda r: (-r["head"], r["total"]))
    else:
        scored.sort(key=lambda r: (-r["head"], -r["total"]))

    return scored, no_score


def _write_cover_sheet(ws, rodeo_name, summary_rows):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 14

    ws["A1"] = rodeo_name
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Final Scoring Report"
    ws["A2"].font = SUBTITLE_FONT
    ws["A3"] = f"Generated {datetime.datetime.now().strftime('%B %d, %Y %I:%M %p')}"
    ws["A3"].font = SUBTITLE_FONT

    header_row = 5
    for col, label in enumerate(("Event", "Type", "Competitors"), start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER

    r = header_row + 1
    for name, type_label, count in summary_rows:
        ws.cell(row=r, column=1, value=name).font = BODY_FONT
        ws.cell(row=r, column=2, value=type_label).font = BODY_FONT
        c = ws.cell(row=r, column=3, value=count)
        c.font = BODY_FONT
        c.alignment = CENTER
        r += 1

    ws.freeze_panes = f"A{header_row + 1}"


def _write_event_sheet(ws, rodeo_name, event, scored, no_score):
    scoring_type = event["scoring_type"]
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 8

    ws["A1"] = rodeo_name
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"{event['name']} \u2014 Final Results"
    ws["A2"].font = SECTION_FONT

    header_row = 4
    score_label = "Time" if scoring.is_timed(scoring_type) else "Score"
    headers = ["Rank", "Competitor", "Hometown", "Sponsor", score_label, "Head"]
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER

    r = header_row + 1
    for i, row in enumerate(scored, start=1):
        full_name, _ = _team_names(row["name"], row["partner"])
        ws.cell(row=r, column=1, value=i).alignment = CENTER
        ws.cell(row=r, column=2, value=full_name)
        ws.cell(row=r, column=3, value=row["hometown"])
        ws.cell(row=r, column=4, value=row["sponsor"])
        score_cell = ws.cell(row=r, column=5, value=scoring.format_number(row["total"], scoring_type))
        score_cell.alignment = CENTER
        head_cell = ws.cell(row=r, column=6, value=row["head"])
        head_cell.alignment = CENTER
        for col in range(1, 7):
            ws.cell(row=r, column=col).font = BODY_FONT
        r += 1

    if not scored:
        ws.cell(row=r, column=1, value="No scored runs yet.").font = BODY_FONT
        r += 1

    if no_score:
        r += 1
        ws.cell(row=r, column=1, value="No Qualifying Score").font = SECTION_FONT
        r += 1
        for row in no_score:
            full_name, _ = _team_names(row["name"], row["partner"])
            statuses = ", ".join(
                scoring.STATUS_LABELS.get(s, s) for s in (row["statuses"] or "").split(",") if s
            )
            ws.cell(row=r, column=2, value=full_name).font = BODY_FONT
            ws.cell(row=r, column=3, value=statuses).font = BODY_FONT
            r += 1

    ws.freeze_panes = f"A{header_row + 1}"


def generate_scoring_report():
    """Builds the full workbook and returns an in-memory BytesIO buffer,
    ready to stream straight to the browser -- nothing is written to
    disk on its own, same as the existing "Download Backup" button."""
    conn = get_conn()
    rodeo_name = get_setting("rodeo_name") or "Rodeo"
    events = conn.execute("SELECT * FROM events ORDER BY position, created_at").fetchall()

    wb = Workbook()
    wb.remove(wb.active)

    used_titles = set()
    summary_rows = []

    for event in events:
        scored, no_score = _event_rows(conn, event)
        summary_rows.append((event["name"], scoring.short_label(event["scoring_type"]), len(scored)))
        ws = wb.create_sheet(_safe_sheet_title(event["name"], used_titles))
        _write_event_sheet(ws, rodeo_name, event, scored, no_score)

    conn.close()

    # Built last (needed the per-event counts first) but displayed
    # first -- index 0 -- since it's the landing page for the workbook.
    cover = wb.create_sheet("Summary", 0)
    _write_cover_sheet(cover, rodeo_name, summary_rows)
    wb.active = 0

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def report_filename(rodeo_name):
    safe_name = re.sub(r"[^A-Za-z0-9_-]+", "_", rodeo_name).strip("_") or "Rodeo"
    stamp = datetime.datetime.now().strftime("%Y-%m-%d")
    return f"{safe_name}_Scoring_Report_{stamp}.xlsx"
